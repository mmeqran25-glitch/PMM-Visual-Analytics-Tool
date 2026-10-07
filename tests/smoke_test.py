from __future__ import annotations

from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import base64
import gzip
import hashlib

import pandas as pd
from openpyxl import load_workbook

from master_utils import load_master_workbook, master_snapshot, validate_master, active_themes, retired_themes, active_dimensions
from share_utils import supervisor_share_html
from snapshot_utils import load_supervisor_snapshot, nodes_of_type, build_supervisor_snapshot
from tree_utils import build_selected_themes_tree_data
from dimension_export import dimension_trace_workbook_bytes
from bi_dashboard import build_study_catalog, filter_catalog, dashboard_counts, novelty_study_rows, context_completeness, _apply_context_display_limit, _top_n_with_other, _country_bi_group, _sector_bi_group
from qualitative_visuals import build_dimension_sankey, build_all_dimensions_sankey, build_theme_study_matrix, build_heatmap_figure, build_cooccurrence_figure, build_wordcloud_frequencies, build_wordcloud_image, build_theme_boundary_profile, build_negative_deviant_case_matrix, build_entreq_audit_table, build_audit_trail_events, build_audit_timeline_figure, ENTREQ_ITEMS
from evidence_matrix import build_evidence_links, build_study_theme_matrix, build_study_dimension_matrix, build_pcl_theme_matrix, build_dimension_evidence_synthesis
from defense_mode import measurement_spec_table, pairwise_boundary_table
from item_pool_prep import item_pool_anchor_table
from item_drafting import build_draft_record, required_facets, _draft_workbook_bytes
from researcher_cache import (
    token_matches,
    save_cached_master,
    load_cached_master,
    clear_cached_master,
)


HEADER_ROW = 3


def _sheet(writer, name, columns, rows=None):
    rows = rows or []
    df = pd.DataFrame(rows, columns=columns)
    df.to_excel(writer, sheet_name=name, index=False, startrow=HEADER_ROW)


def build_workbook(version: str = "v99.123", decision: str = "DEC-999") -> bytes:
    bio = BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as writer:
        readme = pd.DataFrame([
            [f"PMM DIMENSION DERIVATION MASTER - {version} - TEST"],
            ["Synthetic compatibility workbook"],
            [f"CURRENT {decision}: 1/1 studies (100%); Next action: continue review."],
        ])
        readme.to_excel(writer, sheet_name="00_README", index=False, header=False)

        _sheet(
            writer,
            "01_Source_Register",
            [
                "Study_ID", "Title", "Authors", "Year", "Source_Publication_Type",
                "Study_Design", "Methodological_Family", "QA_Judgment", "QA_Percent",
            ],
            [["SR001", "Synthetic Study", "A. Researcher", 2024, "Peer-reviewed journal article", "Survey", "Quantitative", "Include", 0.8]],
        )
        _sheet(
            writer,
            "02_Capability_Eligibility",
            ["Study_ID", "Dimension_Formation_Decision"],
            [["SR001", "Eligible"]],
        )
        _sheet(
            writer,
            "03_Study_Profile",
            [
                "Study_ID", "Title", "Authors", "Year", "Country_Context", "Sector_Context",
                "Organization_Level", "Study_Design", "Methodological_Family",
                "PMM_Model_or_Framework", "Unit_of_Analysis", "Evidence_Source_Role",
                "Study_Family_ID", "Profile_Status",
            ],
            [[
                "SR001", "Synthetic Study", "A. Researcher", 2024, "Yemen",
                "Telecommunications", "Organization", "Cross-sectional survey",
                "Quantitative", "Synthetic PMM", "Organization", "Primary empirical",
                "SF-TEST-01", "Complete",
            ]],
        )
        _sheet(
            writer,
            "04_Verbatim_Evidence",
            [
                "Evidence_ID", "Study_ID", "Meaning_Unit_Verbatim", "Meaning_Unit_Status",
                "PM_Practice_Maturity_Evidence_Gate", "Context_Verbatim", "Original_Author_Term",
                "Author_Parent_Construct", "Author_Defined_Relationship", "Printed_Page",
            ],
            [["EV001", "SR001", "Synthetic evidence", "Approved", "Pass", "Context", "Term", "Parent", "Relation", "1"]],
        )
        _sheet(
            writer,
            "05_First_Order_Coding",
            ["Code_ID", "Evidence_ID", "Study_ID", "First_Order_Code", "Code_Fidelity_Status", "Source_to_Code_Rationale"],
            [["CD-SR001-001", "EV001", "SR001", "Synthetic FOC", "Pass", "Rationale"]],
        )
        _sheet(
            writer,
            "06_DeNovo_Clustering",
            [
                "Mapping_ID", "Code_ID", "Study_ID", "Cluster_ID", "Mapping_Status",
                "Mapping_Rationale", "Boundary_Rationale", "Closest_Competing_Cluster",
            ],
            [["MAP001", "CD-SR001-001", "SR001", "PCL-001", "Stable", "Rationale", "", ""]],
        )
        _sheet(
            writer,
            "06A_Cluster_Register",
            [
                "Cluster_ID", "Working_Cluster_Label", "Cluster_Status", "Core_Organizational_Function",
                "Operational_Definition", "Inclusion_Boundary", "Exclusion_Boundary", "Nearest_Conceptual_Neighbours",
            ],
            [["PCL-001", "Synthetic Cluster", "Stable", "Function", "Definition", "Include", "Exclude", "Neighbour"]],
        )
        _sheet(
            writer,
            "07_Descriptive_Themes",
            [
                "Theme_ID", "Working_Theme_Label", "Included_Cluster_IDs", "Theme_Status",
                "Central_Organizing_Concept", "Theme_Boundary", "Closest_Competing_Theme",
            ],
            [
                ["THM-001", "Synthetic Theme", "PCL-001", "Stable", "Concept", "Boundary", ""],
                ["THM-016", "Embedding sustainability across project-management systems and stakeholder/supply-chain interfaces", "PCL-001", "Retired – DEC-351 functional-boundary refinement", "Historical concept", "Historical boundary", ""],
            ],
        )
        _sheet(
            writer,
            "08_Candidate_Dimensions",
            [
                "Dimension_ID", "Candidate_Dimension_Name", "Supporting_Theme_IDs", "Dimension_Status",
                "Analytical_Definition", "Core_Capability_Logic", "Evidence_Breadth", "Cross_Model_Support",
                "Cross_Context_Support", "Measurement_Evidence_Link", "Contradictory_or_Boundary_Evidence", "Notes",
            ],
            [["DIM-001", "Synthetic Dimension", "THM-001", "Stable", "Definition", "Logic", "Breadth", "", "", "", "", ""]],
        )
        dim_ws = writer.sheets["08_Candidate_Dimensions"]
        pair_headers = [
            "Audit_ID", "Dimension_A", "Dimension_B", "Shared_Source_Records",
            "Boundary_Risk", "Controlling_Discriminator", "Final_Decision",
            "Pairwise_Rationale", "Measurement_Caution",
        ]
        pair_values = [
            "D588-01", "DIM-001", "DIM-002", 1, "Low",
            "Synthetic discriminator", "PASS – RETAIN DISTINCT",
            "Synthetic pairwise rationale", "Synthetic measurement caution",
        ]
        for col_idx, value in enumerate(pair_headers, start=1):
            dim_ws.cell(row=20, column=col_idx, value=value)
        for col_idx, value in enumerate(pair_values, start=1):
            dim_ws.cell(row=21, column=col_idx, value=value)

        measurement_headers = [
            "Dimension_ID", "Pre_Item_Measurement_Form", "Required_Content_Facets",
            "Item_Referent", "Exclude_or_Contamination_Rule", "Applicability_NA_Control",
            "Aggregation_Control", "Validation_Gate", "Pre_Item_Disposition",
        ]
        measurement_values = [
            "DIM-001", "Synthetic multifacet content-composite",
            "Facet A; Facet B", "Synthetic organizational PM practice",
            "Exclude outcomes", "Use N/A where not observable",
            "Preserve facet coverage", "Expert CVI + pilot + empirical testing",
            "LOCK PRE-ITEM ARCHITECTURE; ITEM POOL NEXT",
        ]
        for col_idx, value in enumerate(measurement_headers, start=1):
            dim_ws.cell(row=25, column=col_idx, value=value)
        for col_idx, value in enumerate(measurement_values, start=1):
            dim_ws.cell(row=26, column=col_idx, value=value)
        _sheet(
            writer,
            "09_Novelty_Tracking",
            [
                "Novelty_Record_ID", "Batch_or_Sequence", "Study_ID", "Extraction_Order",
                "Eligible_Original_Terms_Count", "New_First_Order_Code?", "New_Core_Capability?",
                "New_Cluster?", "Boundary_Change?", "Split_or_Merge_Triggered?",
                "Mainly_Reinforcement?", "Novelty_Summary", "Coverage_Gap_Addressed",
                "Cumulative_Stability_Observation", "Decision_for_Next_Sampling",
            ],
            [[
                "NOV-001", "Synthetic expansion", "SR001", 1, 1, "Yes – 1 new FOC",
                "No", "Yes – PCL-001 Provisional", "No", "No",
                "Mostly reinforcement", "Synthetic novelty record", "Synthetic gap",
                "Initial stability observation", "Continue",
            ]],
        )
        _sheet(writer, "10_Study_Families", ["Study_ID"], [["SR001"]])
        _sheet(writer, "11_Decision_Log", ["Decision_ID"], [[decision]])

    return bio.getvalue()


def main():
    # Academic branding is part of the public/researcher interface contract.
    app_source = Path("app.py").read_text(encoding="utf-8")
    assert "رسالة ماجستير للباحث – معاذ عبدالقوي عباس مقران" in app_source
    assert "جامعة صنعاء · كلية الهندسة · الدراسات العليا" in app_source
    assert 'page_icon="🎓"' in app_source
    assert "render_academic_brand_header" in app_source
    assert "render_sidebar_identity" in app_source
    assert "render_academic_footer" in app_source
    assert "brand-logo-tile" in app_source
    assert "padding-top:2.85rem !important" in app_source
    assert "_local_image_data_uri" in app_source
    assert 'ASSET_DIR / "sanaa_university_logo.jpg"' in app_source
    assert 'ASSET_DIR / "faculty_engineering_logo.jpg"' in app_source
    assert (Path("assets") / "sanaa_university_logo.jpg").exists()
    assert (Path("assets") / "faculty_engineering_logo.jpg").exists()
    from PIL import Image
    logo = Image.open(Path("assets") / "faculty_engineering_logo.jpg")
    assert logo.width >= 150 and logo.height >= 210

    arbitrary_filename = "THIS_NAME_CAN_CHANGE_EVERY_DAY_v2045.xlsx"
    payload = build_workbook()

    frames, structure = load_master_workbook(payload)
    assert validate_master(frames) == [], validate_master(frames)

    links = build_evidence_links(frames)
    assert len(links) == 1, links
    assert links.iloc[0]["Study_ID"] == "SR001", links
    assert links.iloc[0]["Theme_ID"] == "THM-001", links
    assert links.iloc[0]["Dimension_ID"] == "DIM-001", links
    assert links.iloc[0]["Evidence_Family_ID"] == "SF-TEST-01", links

    # Embedded DEC-588 / DEC-589 tables must be recoverable from the existing 08 sheet.
    pairwise = pairwise_boundary_table(frames)
    assert len(pairwise) == 1, pairwise
    assert pairwise.iloc[0]["Audit_ID"] == "D588-01", pairwise
    assert pairwise.iloc[0]["Final_Decision"] == "PASS – RETAIN DISTINCT", pairwise

    measurement = measurement_spec_table(frames)
    assert len(measurement) == 1, measurement
    assert measurement.iloc[0]["Dimension_ID"] == "DIM-001", measurement
    assert measurement.iloc[0]["Required_Content_Facets"] == "Facet A; Facet B", measurement
    assert "ITEM POOL NEXT" in measurement.iloc[0]["Pre_Item_Disposition"], measurement

    anchors = item_pool_anchor_table(frames, "DIM-001")
    assert len(anchors) == 1, anchors
    assert anchors.iloc[0]["Code_ID"] == "CD-SR001-001", anchors
    assert anchors.iloc[0]["Theme_ID"] == "THM-001", anchors
    assert anchors.iloc[0]["Cluster_ID"] == "PCL-001", anchors
    assert anchors.iloc[0]["Original_Author_Term"] == "Term", anchors

    # Item drafting must preserve locked source traceability and produce an expert-review package.
    facets = required_facets(frames, "DIM-001")
    assert facets == ["Facet A", "Facet B"], facets
    draft = build_draft_record(frames, "DIM-001", facets[0], anchors, 1)
    assert draft["Draft_Item_ID"] == "DRAFT-DIM-001-001", draft
    assert draft["Supporting_FOCs"] == "CD-SR001-001", draft
    assert draft["Supporting_PCLs"] == "PCL-001", draft
    assert draft["Supporting_Studies"] == "SR001", draft
    assert draft["Required_Content_Facet"] == "Facet A", draft
    assert draft["Exclude_or_Contamination_Rule"] == "Exclude outcomes", draft

    draft_df = pd.DataFrame([draft])
    draft_df.loc[0, "Item_Concept"] = "Synthetic item concept"
    draft_df.loc[0, "Candidate_Item_EN"] = "Our organization uses the synthetic PM practice."
    package_bytes = _draft_workbook_bytes(draft_df, measurement.iloc[0], 5)
    package_wb = load_workbook(BytesIO(package_bytes), read_only=True, data_only=False)
    assert package_wb.sheetnames == [
        "00_Instructions",
        "01_Draft_Items",
        "02_DEC589_Spec",
        "03_Expert_Ratings",
        "04_CVI_Summary",
    ]
    expert_ws = package_wb["03_Expert_Ratings"]
    assert expert_ws.max_row == 6  # header + 5 expert-rating rows

    stm = build_study_theme_matrix(links, mode="Presence", theme_ids=["THM-001"])
    assert stm.loc["SR001", "THM-001"] == "●", stm
    sdm = build_study_dimension_matrix(links, mode="FOC Count", dimension_ids=["DIM-001"])
    assert int(sdm.loc["SR001", "DIM-001"]) == 1, sdm
    ptm = build_pcl_theme_matrix(frames, ["THM-001"])
    assert ptm.loc["PCL-001", "THM-001"] == "●", ptm

    synth = build_dimension_evidence_synthesis(links, "DIM-001")
    assert len(synth) == 1, synth
    assert synth.iloc[0]["Study_ID"] == "SR001", synth
    assert synth.iloc[0]["Factors_or_Contributions"] == "Synthetic Cluster", synth
    assert int(synth.iloc[0]["FOC_Count"]) == 1, synth

    snap = master_snapshot(frames)
    assert snap["version"] == "v99.123", snap
    assert snap["decision"] == "DEC-999", snap
    assert snap["active_focs"] == 1, snap
    assert snap["mapped_focs"] == 1, snap
    assert snap["active_clusters"] == 1, snap
    assert snap["active_themes"] == 1, snap
    assert snap["active_dimensions"] == 1, snap

    page = supervisor_share_html(frames, snap)
    assert "Synthetic Dimension" in page
    assert "PCL-001" in page
    assert "Synthetic FOC" in page
    assert "Synthetic evidence" in page

    # The presentation generator never receives a workbook filename.
    assert arbitrary_filename not in page
    assert ".xlsx" not in page.lower()

    # Retired Themes remain excluded from current-state logic but can be opened explicitly for audit history.
    current_themes = active_themes(frames)
    history_themes = retired_themes(frames)
    assert current_themes["Theme_ID"].astype(str).tolist() == ["THM-001"]
    assert history_themes["Theme_ID"].astype(str).tolist() == ["THM-016"]

    current_only_tree = build_selected_themes_tree_data(frames, ["THM-016"], include_retired=False)
    assert len(current_only_tree["children"]) == 0

    audit_tree = build_selected_themes_tree_data(frames, ["THM-016"], include_retired=True)
    assert len(audit_tree["children"]) == 1
    retired_node = audit_tree["children"][0]
    assert retired_node["id"] == "THM-016"
    assert "retired" in retired_node["status"].lower()
    assert len(nodes_of_type(retired_node, "cluster")) == 1
    assert len(nodes_of_type(retired_node, "code")) == 0

    live_supervisor = build_supervisor_snapshot(frames)
    assert len(live_supervisor.get("retired_themes", [])) == 1
    assert live_supervisor["retired_themes"][0]["id"] == "THM-016"

    # Current-dimension selection must ignore retained historical/suspended rows
    # and prefer a later audited working set when canonical current statuses are absent.
    dim_test = pd.DataFrame([
        {
            "Dimension_ID": "DIM-001",
            "Candidate_Dimension_Name": "Old historical label",
            "Supporting_Theme_IDs": "THM-001; THM-002",
            "Underlying_Cluster_IDs": "PCL-001; PCL-002",
            "Dimension_Status": "Historical / Superseded – excluded from fresh SG4 derivation",
            "Analytical_Definition": "Old definition",
            "Core_Capability_Logic": "Old logic",
        },
        {
            "Dimension_ID": "DIM-007",
            "Candidate_Dimension_Name": "Suspended exploratory candidate",
            "Supporting_Theme_IDs": "THM-003",
            "Underlying_Cluster_IDs": "PCL-003",
            "Dimension_Status": "SUSPENDED – SG4 FROZEN PENDING SG2/SG3 BOUNDARY REPAIR",
            "Analytical_Definition": "Suspended definition",
            "Core_Capability_Logic": "Suspended",
        },
        {
            "Dimension_ID": "DIM-001",
            "Candidate_Dimension_Name": "Current working governance dimension",
            "Supporting_Theme_IDs": "Strategic alignment; authorization; priority decisions",
            "Underlying_Cluster_IDs": "Not listed in the working audit block",
            "Dimension_Status": "",
            "Analytical_Definition": "Set organizational direction and formal project-choice decisions.",
            "Core_Capability_Logic": "PASS – PROVISIONAL",
        },
        {
            "Dimension_ID": "DIM-001",
            "Candidate_Dimension_Name": "Boundary challenge narrative only",
            "Supporting_Theme_IDs": "",
            "Underlying_Cluster_IDs": "",
            "Dimension_Status": "",
            "Analytical_Definition": "",
            "Core_Capability_Logic": "",
        },
    ])
    dim_frames = {"08_Candidate_Dimensions": dim_test}
    current_dims = active_dimensions(dim_frames)
    assert len(current_dims) == 1
    assert current_dims.iloc[0]["Candidate_Dimension_Name"] == "Current working governance dimension"
    assert current_dims.iloc[0]["Dimension_Status"] == "PASS – PROVISIONAL"
    assert current_dims.iloc[0]["Supporting_Theme_IDs"] == "THM-001; THM-002"
    assert current_dims.iloc[0]["Underlying_Cluster_IDs"] == "PCL-001; PCL-002"

    # Research BI catalog keeps study universes and study-level metadata distinct.
    bi_catalog = build_study_catalog(frames)
    assert len(bi_catalog) == 1
    bi_row = bi_catalog.iloc[0]
    assert bi_row["Study_ID"] == "SR001"
    assert int(bi_row["Year"]) == 2024
    assert bi_row["Country_Context"] == "Yemen"
    assert bi_row["Sector_Context"] == "Telecommunications"
    assert bi_row["Methodology_Group"] == "Quantitative"
    assert bi_row["Publication_Group"] == "Journal Article"
    assert bi_row["Eligibility_Group"] == "Eligible"
    assert bi_row["Evidence_Records"] == 1
    assert bi_row["Pass_Evidence"] == 1
    assert bi_row["FOC_Count"] == 1
    assert round(float(bi_row["QA_Percent_Display"]), 1) == 80.0

    assert len(filter_catalog(bi_catalog, "All Sources")) == 1
    assert len(filter_catalog(bi_catalog, "Profiled Studies")) == 1
    assert len(filter_catalog(bi_catalog, "Evidence Contributors")) == 1
    assert len(filter_catalog(bi_catalog, "Pass Contributors")) == 1

    bi_counts = dashboard_counts(bi_catalog, frames)
    assert bi_counts["sources"] == 1
    assert bi_counts["profiled"] == 1
    assert bi_counts["contributors"] == 1
    assert bi_counts["pass_contributors"] == 1
    assert bi_counts["evidence_records"] == 1
    assert bi_counts["pass_evidence"] == 1
    assert bi_counts["period"] == "2024–2024"

    completeness = context_completeness(bi_catalog)
    assert not completeness.empty
    assert float(completeness.loc[completeness["Field"] == "Country_Context", "Percent"].iloc[0]) == 100.0
    assert float(completeness.loc[completeness["Field"] == "Sector_Context", "Percent"].iloc[0]) == 100.0

    assert _country_bi_group("South Africa – national government / South African Public Service") == "South Africa"
    assert _country_bi_group("No empirical country context specified; conceptual article.") == "Not specified / Conceptual"
    assert _country_bi_group("Austria and Hungary / Austrian-Hungarian cross-border region") == "International / Multi-country"
    assert _sector_bi_group("Telecommunications / telecom industry") == "Telecommunications / ICT"
    assert _sector_bi_group("Solar energy / Engineering, Procurement and Construction (EPC)") == "Construction / Engineering / EPC"

    topn_input = pd.DataFrame({
        "Display_Category": ["A", "B", "C", "D"],
        "Studies": [10, 8, 6, 4],
    })
    topn = _top_n_with_other(topn_input, "Display_Category", "Studies", 2)
    assert topn["Display_Category"].tolist() == ["A", "B", "Other"]
    assert int(topn.loc[topn["Display_Category"] == "Other", "Studies"].iloc[0]) == 10
    assert int(topn["Studies"].sum()) == int(topn_input["Studies"].sum())

    context_counts = pd.DataFrame({
        "Country_Context": ["A", "B", "C", "D"],
        "Studies": [10, 8, 6, 4],
    })
    assert len(_apply_context_display_limit(context_counts, "All")) == 4
    assert len(_apply_context_display_limit(context_counts, "Top 10")) == 4
    assert len(_apply_context_display_limit(context_counts, "Top 2")) == 2

    novelty = novelty_study_rows(frames, ["SR001"])
    assert len(novelty) == 1
    assert bool(novelty.iloc[0]["New_FOC_Flag"])
    assert bool(novelty.iloc[0]["New_Cluster_Flag"])
    assert not bool(novelty.iloc[0]["Boundary_Change_Flag"])
    assert bool(novelty.iloc[0]["Reinforcement_Flag"])

    # Qualitative thesis visuals resolve the current derivation structure and study support.
    word_freq, word_records = build_wordcloud_frequencies(frames, "First-Order Code labels", max_words=50)
    assert isinstance(word_freq, dict)
    assert len(word_freq) >= 1
    assert len(word_records) >= 1
    word_png = build_wordcloud_image(word_freq)
    assert word_png[:8] == b"\x89PNG\r\n\x1a\n"
    assert len(word_png) > 1000

    entreq = build_entreq_audit_table(frames)
    assert len(ENTREQ_ITEMS) == 21
    assert len(entreq) == 21
    assert entreq["No."].tolist() == list(range(1, 22))
    assert set(entreq["Domain"]) == {
        "Introduction",
        "Methods & methodology",
        "Literature search & selection",
        "Appraisal",
        "Synthesis of findings",
    }
    assert set(entreq["Status"]) == {"Not assessed"}
    study_char = entreq.loc[entreq["No."].eq(8), "Project evidence hint"].iloc[0]
    assert study_char == "Strong project hint"
    coding_hint = entreq.loc[entreq["No."].eq(17), "Project evidence hint"].iloc[0]
    assert coding_hint == "Strong project hint"

    assert "Provisional review" in entreq.columns
    assert "Review rationale / next action" in entreq.columns
    assert entreq.loc[entreq["No."].eq(10), "Provisional review"].iloc[0] == "Pending appraisal phase"
    assert entreq.loc[entreq["No."].eq(17), "Provisional review"].iloc[0] == "Ready – strong evidence"

    timeline_frames = dict(frames)
    timeline_frames["11_Decision_Log"] = pd.DataFrame([
        {
            "Decision_ID": "DEC-401",
            "Decision_Title": "Boundary review",
            "Decision_Summary": "Boundary review for PCL-001 and THM-001",
            "Decision_Rationale": "External heterogeneity check",
            "Stage": "SG3",
        },
        {
            "Decision_ID": "DEC-402",
            "Decision_Title": "Cluster split",
            "Decision_Summary": "Split PCL-001 into two clusters",
            "Decision_Rationale": "Internal heterogeneity",
            "Stage": "SG2",
        },
        {
            "Decision_ID": "DEC-403",
            "Decision_Title": "FOC reassignment",
            "Decision_Summary": "Reassigned CD-SR001-001 to PCL-002",
            "Decision_Rationale": "Closer functional fit",
            "Stage": "SG2",
        },
        {
            "Decision_ID": "DEC-404",
            "Decision_Title": "Theme retirement",
            "Decision_Summary": "Retired THM-016 after functional-boundary refinement",
            "Decision_Rationale": "Superseded structure",
            "Stage": "SG3",
        },
        {
            "Decision_ID": "DEC-405",
            "Decision_Title": "Theme merge",
            "Decision_Summary": "Merged two overlapping Themes",
            "Decision_Rationale": "Redundant organizing concept",
            "Stage": "SG3",
        },
    ])
    audit_events = build_audit_trail_events(timeline_frames)
    assert not audit_events.empty
    assert audit_events["DEC_Sequence"].min() <= 401
    assert {"Boundary Review", "Split", "Reassignment", "Retirement / Withdrawal", "Merge"}.issubset(set(audit_events["Event_Type"]))
    audit_fig = build_audit_timeline_figure(audit_events)
    assert len(audit_fig.data) >= 1

    boundary_profile = build_theme_boundary_profile(frames, "THM-001")
    assert boundary_profile["Theme_ID"] == "THM-001"
    assert boundary_profile["Clusters"] == 1
    assert boundary_profile["FOCs"] == 1
    assert boundary_profile["Studies"] == 1
    assert boundary_profile["Central_Organizing_Concept"] == "Concept"
    assert boundary_profile["Theme_Boundary"] == "Boundary"
    assert not boundary_profile["Cluster_Table"].empty
    assert boundary_profile["Boundary_Completeness"]["Operational_Definition"] == 100.0
    assert boundary_profile["Boundary_Completeness"]["Inclusion_Boundary"] == 100.0
    assert boundary_profile["Boundary_Completeness"]["Exclusion_Boundary"] == 100.0

    negative_matrix = build_negative_deviant_case_matrix(frames)
    assert not negative_matrix.empty
    retired_rows = negative_matrix[negative_matrix["Case_Type"].eq("Retired Theme")]
    assert "THM-016" in retired_rows["Case_ID"].tolist()
    assert retired_rows["Current_or_Historical"].eq("Historical").all()

    all_sankey = build_all_dimensions_sankey(frames)
    assert len(all_sankey.data) == 1
    assert all_sankey.data[0].type == "sankey"
    all_labels = [str(x) for x in all_sankey.data[0].node.label]
    assert any("DIM-001" in x for x in all_labels)
    assert any("THM-001" in x for x in all_labels)

    sankey = build_dimension_sankey(frames, "DIM-001")
    assert len(sankey.data) == 1
    assert sankey.data[0].type == "sankey"
    assert len(sankey.data[0].link.source) >= 2

    qmatrix = build_theme_study_matrix(frames, "Evidence count")
    assert not qmatrix.empty
    assert ("THM-001", "Synthetic Theme") in qmatrix.index
    assert "SR001" in qmatrix.columns
    assert int(qmatrix.loc[("THM-001", "Synthetic Theme"), "SR001"]) == 1
    assert not any(idx[0] == "THM-016" for idx in qmatrix.index)

    heatmap, shown = build_heatmap_figure(qmatrix, 15)
    assert len(heatmap.data) == 1
    assert heatmap.data[0].type == "heatmap"
    assert not shown.empty

    network, edge_table = build_cooccurrence_figure(frames, 1)
    assert len(network.data) >= 1
    assert network.data[-1].type == "scatter"
    assert edge_table.empty  # only one active Theme in the synthetic workbook

    # Selected Dimension can be exported as a complete Excel traceability package.
    export_bytes, export_summary = dimension_trace_workbook_bytes(frames, "DIM-001")
    assert export_summary["Dimension_ID"] == "DIM-001"
    assert export_summary["Theme_Count"] == 1
    assert export_summary["Cluster_Count"] == 1
    assert export_summary["FOC_Count"] == 1
    assert export_summary["Evidence_Count"] == 1
    assert export_summary["Study_Count"] == 1

    export_wb = load_workbook(BytesIO(export_bytes), read_only=True, data_only=False)
    assert export_wb.sheetnames == [
        "00_Export_Info",
        "01_Full_Trace",
        "02_Themes",
        "03_Clusters",
        "04_FOCs",
        "05_Studies",
        "06_Tree_Index",
    ]
    trace_ws = export_wb["01_Full_Trace"]
    headers = [cell.value for cell in next(trace_ws.iter_rows(min_row=1, max_row=1))]
    first_row = [cell.value for cell in next(trace_ws.iter_rows(min_row=2, max_row=2))]
    trace = dict(zip(headers, first_row))
    assert trace["Dimension_ID"] == "DIM-001"
    assert trace["Theme_ID"] == "THM-001"
    assert trace["Cluster_ID"] == "PCL-001"
    assert trace["Code_ID"] == "CD-SR001-001"
    assert trace["Evidence_ID"] == "EV001"
    assert trace["Study_ID"] == "SR001"
    assert trace["Meaning_Unit_Verbatim"] == "Synthetic evidence"

    tree_ws = export_wb["06_Tree_Index"]
    tree_values = list(tree_ws.iter_rows(min_row=2, values_only=True))
    tree_types = {row[1] for row in tree_values}
    assert {"Dimension", "Theme", "Cluster", "First-Order Code", "Meaning Unit", "Study"}.issubset(tree_types)

    # Published supervisor snapshot must load from the repository without any Excel file.
    published = load_supervisor_snapshot()
    assert published is not None
    tree = published["tree"]
    assert len(nodes_of_type(tree, "dimension")) == 4
    assert len(nodes_of_type(tree, "theme")) == 13
    assert len(nodes_of_type(tree, "cluster")) == 77
    assert len(nodes_of_type(tree, "code")) == 1277
    assert len(nodes_of_type(tree, "evidence")) == 0
    assert len(nodes_of_type(tree, "study")) == 0

    # The committed compact payload itself must not contain source-near fields or workbook references.
    parts = sorted(Path("supervisor_snapshot_chunks").glob("part*.b64"))
    assert parts
    encoded = "".join(p.read_text(encoding="ascii").strip() for p in parts)
    compact_text = gzip.decompress(base64.b64decode(encoded)).decode("utf-8")
    for forbidden in [".xlsx", "Meaning_Unit_Verbatim", "Context_Verbatim", "Study_ID", "Evidence_ID"]:
        assert forbidden not in compact_text, forbidden

    # Researcher workbook cache persists independently of a Streamlit browser session.
    with TemporaryDirectory() as td:
        saved = save_cached_master("changing_name_v99.xlsx", payload, cache_root=td)
        assert saved["filename"] == "changing_name_v99.xlsx"
        cached = load_cached_master(cache_root=td)
        assert cached is not None
        cached_name, cached_bytes, cached_meta = cached
        assert cached_name == "changing_name_v99.xlsx"
        assert cached_bytes == payload
        assert cached_meta["sha256"] == hashlib.sha256(payload).hexdigest()
        clear_cached_master(cache_root=td)
        assert load_cached_master(cache_root=td) is None

    # Token verification is hash-based; tests never require the production raw researcher key.
    unit_secret = "unit-test-secret"
    unit_hash = hashlib.sha256(unit_secret.encode("utf-8")).hexdigest()
    assert token_matches(unit_secret, expected_hash=unit_hash)
    assert not token_matches("wrong-secret", expected_hash=unit_hash)

    # Workbook evolution is data-driven: changing metadata does not require code changes.
    frames2, _ = load_master_workbook(build_workbook(version="v2045.7", decision="DEC-2045"))
    snap2 = master_snapshot(frames2)
    assert snap2["version"] == "v2045.7", snap2
    assert snap2["decision"] == "DEC-2045", snap2

    print("SMOKE TEST PASSED")
    print(f"Sheets loaded: {len(structure)}")
    print(f"Snapshot: {snap}")
    print(f"Future snapshot: {snap2}")


if __name__ == "__main__":
    main()


def test_qualitative_visual_version_marker():
    import qualitative_visuals
    assert qualitative_visuals.QUAL_VIS_VERSION == "v0.14.3"
