from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook


DEFAULT_LOCAL_DATA_DIR = Path(
    r"E:\مراجع للرسالة ان شاء الله\بداية تجهيز الرساله\قاعدة بيانات منصةرسالة الماجسيتر الخاص بمعاذ"
)
LOCAL_DATA_DIR_ENV = "PMM_DATA_DIR"

MASTER_SIGNATURE_SHEETS = {
    "01_Source_Register",
    "02_Capability_Eligibility",
    "04_Verbatim_Evidence",
    "05_First_Order_Coding",
    "06_DeNovo_Clustering",
    "06A_Cluster_Register",
    "07_Descriptive_Themes",
    "08_Candidate_Dimensions",
    "11_Decision_Log",
}
AUDIT_ARCHIVE_SIGNATURE = "00_AUDIT_ARCHIVE_INDEX"
PRISMA_SIGNATURE_SHEETS = {"Master Screening", "Search Log"}


def configured_local_data_dir() -> Path:
    """Return the configured local Excel database directory.

    The environment variable allows the same repository to run on another PC
    without changing source code. When it is absent, Moaz's Windows E: path is
    used as the local default.
    """
    raw = os.environ.get(LOCAL_DATA_DIR_ENV, "").strip()
    return Path(raw) if raw else DEFAULT_LOCAL_DATA_DIR


def _safe_sheet_names(path: Path) -> List[str]:
    try:
        wb = load_workbook(path, read_only=True, data_only=False)
        names = list(wb.sheetnames)
        wb.close()
        return names
    except Exception:
        return []


def classify_workbook(path: Path) -> str:
    names = set(_safe_sheet_names(path))
    if not names:
        return "unreadable"
    if AUDIT_ARCHIVE_SIGNATURE in names:
        return "audit_archive"
    if PRISMA_SIGNATURE_SHEETS.issubset(names):
        return "prisma"
    if MASTER_SIGNATURE_SHEETS.issubset(names):
        return "master"
    return "other"


def _version_key(name: str) -> Tuple[int, Tuple[int, ...]]:
    """Extract DEC and v-style version numbers for semantic newest-first sorting."""
    decs = [int(x) for x in re.findall(r"DEC[-_ ]?(\d+)", name, flags=re.IGNORECASE)]
    dec = max(decs) if decs else -1

    versions = []
    for raw in re.findall(r"(?:^|[_ -])v(\d+(?:\.\d+)*(?:[A-Za-z])?)", name, flags=re.IGNORECASE):
        nums = tuple(int(x) for x in re.findall(r"\d+", raw))
        if nums:
            versions.append(nums)
    version = max(versions) if versions else tuple()
    return dec, version


def _candidate_sort_key(path: Path):
    dec, version = _version_key(path.name)
    try:
        mtime = path.stat().st_mtime
    except OSError:
        mtime = 0.0
    return (dec, version, mtime, path.name.lower())


def discover_local_excel_sources(
    data_dir: Path | str | None = None,
    recursive: bool = False,
) -> Dict[str, object]:
    """Discover and classify local Excel workbooks without modifying them.

    Returns latest candidates for the analytical MASTER, historical Audit
    Archive, and PRISMA screening workbook plus a classified inventory.
    """
    root = Path(data_dir) if data_dir is not None else configured_local_data_dir()
    result: Dict[str, object] = {
        "root": root,
        "available": False,
        "master": None,
        "audit_archive": None,
        "prisma": None,
        "inventory": [],
        "error": None,
    }

    if not root.exists() or not root.is_dir():
        result["error"] = "Configured folder is not accessible from the machine running Streamlit."
        return result

    result["available"] = True
    iterator = root.rglob("*.xlsx") if recursive else root.glob("*.xlsx")
    files = [
        p for p in iterator
        if p.is_file()
        and not p.name.startswith("~$")
        and not p.name.startswith(".")
    ]

    buckets: Dict[str, List[Path]] = {
        "master": [],
        "audit_archive": [],
        "prisma": [],
        "other": [],
        "unreadable": [],
    }

    inventory = []
    for path in files:
        kind = classify_workbook(path)
        buckets.setdefault(kind, []).append(path)
        try:
            modified = path.stat().st_mtime
            size = path.stat().st_size
        except OSError:
            modified = 0.0
            size = 0
        inventory.append({
            "Filename": path.name,
            "Type": kind,
            "Path": str(path),
            "Size_Bytes": size,
            "Modified_Timestamp": modified,
        })

    for key in ["master", "audit_archive", "prisma"]:
        if buckets[key]:
            result[key] = sorted(buckets[key], key=_candidate_sort_key, reverse=True)[0]

    result["inventory"] = sorted(
        inventory,
        key=lambda x: (x["Type"], x["Modified_Timestamp"], x["Filename"]),
        reverse=True,
    )
    return result


def read_local_workbook(path: Path | str | None) -> Optional[Tuple[str, bytes, dict]]:
    if path is None:
        return None
    p = Path(path)
    try:
        data = p.read_bytes()
        stat = p.stat()
    except Exception:
        return None
    return (
        p.name,
        data,
        {
            "path": str(p),
            "size_bytes": len(data),
            "mtime": stat.st_mtime,
            "source": "local-folder",
        },
    )


def local_source_status(source: Dict[str, object]) -> Dict[str, str]:
    """Compact human-readable source status for UI/tests."""
    return {
        "root": str(source.get("root") or ""),
        "available": "yes" if source.get("available") else "no",
        "master": Path(source["master"]).name if source.get("master") else "",
        "audit_archive": Path(source["audit_archive"]).name if source.get("audit_archive") else "",
        "prisma": Path(source["prisma"]).name if source.get("prisma") else "",
    }
