from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any, Dict, Iterable, List, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from master_utils import (
    active_cluster_register,
    active_dimensions,
    active_themes,
    cluster_members,
    master_snapshot,
    split_ids,
)


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
TITLE_FILL = PatternFill("solid", fgColor="0F6B78")
SUBTITLE_FILL = PatternFill("solid", fgColor="D9EAF0")
CONTROL_FILL = PatternFill("solid", fgColor="E4DFEC")
CAUTION_FILL = PatternFill("solid", fgColor="FCE4D6")
WHITE_FONT = Font(color="FFFFFF", bold=True)
HEADER_FONT = Font(color="FFFFFF", bold=True)
LINKED_FONT = Font(color="008000")
STATIC_FONT = Font(color="666666")
CONTROL_FONT = Font(color="7030A0", bold=True)
THIN_GRAY = Side(style="thin", color="D9E1F2")


def _txt(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    text = str(value)
    text = ILLEGAL_CHARACTERS_RE.sub("", text)
    return text[:32767]


def _first(row: Dict[str, Any], names: Iterable[str]) -> str:
    for name in names:
        value = row.get(name)
        if _txt(value).strip():
            return _txt(value).strip()
    return ""


def _study_lookup(frames: Dict[str, pd.DataFrame]) -> Dict[str, Dict[str, Any]]:
    out: Dict[str, Dict[str, Any]] = {}
    for sheet in ["01_Source_Register", "03_Study_Profile"]:
        df = frames.get(sheet, pd.DataFrame())
        if df.empty or "Study_ID" not in df.columns:
            continue
        for _, row in df.iterrows():
            sid = _txt(row.get("Study_ID")).strip()
            if not sid:
                continue
            target = out.setdefault(sid, {})
            for col, value in row.to_dict().items():
                if col not in target or not _txt(target.get(col)).strip():
                    target[col] = value
    return out


def _study_fields(meta: Dict[str, Any]) -> Dict[str, str]:
    return {
        "Study_Title": _first(meta, ["Study_Title", "Title", "Source_Title"]),
        "Authors": _first(meta, ["Authors", "Author", "Study_Authors"]),
        "Publication_Year": _first(meta, ["Publication_Year", "Year", "Study_Year"]),
        "DOI": _first(meta, ["DOI", "Digital_Object_Identifier"]),
        "Country_Context": _first(meta, ["Country_Context", "Country"]),
        "Sector_Context": _first(meta, ["Sector_Context", "Sector"]),
        "Methodological_Family": _first(meta, ["Methodological_Family", "Methodology", "Method"]),
        "Publication_Type": _first(meta, ["Source_Publication_Type", "Publication_Type", "Source_Type"]),
    }


def build_dimension_trace_tables(
    frames: Dict[str, pd.DataFrame],
    dimension_id: str,
) -> Tuple[Dict[str, Any], Dict[str, pd.DataFrame]]:
    dims = active_dimensions(frames)
    themes = active_themes(frames)
    clusters = active_cluster_register(frames)

    did = str(dimension_id).strip()
    dmatch = dims[dims.get("Dimension_ID", pd.Series(dtype=str)).astype(str).str.strip().eq(did)]
    if dmatch.empty:
        raise ValueError(f"Current dimension not found: {did}")
    drow = dmatch.iloc[0]

    theme_ids = split_ids(drow.get("Supporting_Theme_IDs"), "THM")
    theme_lookup = {
        str(r.get("Theme_ID", "")).strip(): r
        for _, r in themes.iterrows()
        if str(r.get("Theme_ID", "")).strip()
    }
    cluster_lookup = {
        str(r.get("Cluster_ID", "")).strip(): r
        for _, r in clusters.iterrows()
        if str(r.get("Cluster_ID", "")).strip()
    }
    studies = _study_lookup(frames)

    full_rows: List[Dict[str, Any]] = []
    tree_rows: List[Dict[str, Any]] = []
    theme_rows: List[Dict[str, Any]] = []
    cluster_rows: List[Dict[str, Any]] = []

    dimension_name = _txt(drow.get("Candidate_Dimension_Name"))
    dimension_status = _txt(drow.get("Dimension_Status"))
    analytical_definition = _txt(drow.get("Analytical_Definition"))

    tree_rows.append({
        "Level": 1,
        "Node_Type": "Dimension",
        "Node_ID": did,
        "Parent_ID": "",
        "Label": dimension_name,
        "Status": dimension_status,
        "Path": did,
    })

    for tid in theme_ids:
        tr = theme_lookup.get(tid)
        if tr is None:
            continue

        theme_label = _txt(tr.get("Working_Theme_Label"))
        theme_status = _txt(tr.get("Theme_Status"))
        cids = split_ids(tr.get("Included_Cluster_IDs"), "PCL")

        tree_rows.append({
            "Level": 2,
            "Node_Type": "Theme",
            "Node_ID": tid,
            "Parent_ID": did,
            "Label": theme_label,
            "Status": theme_status,
            "Path": f"{did} > {tid}",
        })

        theme_rows.append({
            "Dimension_ID": did,
            "Dimension_Name": dimension_name,
            "Theme_ID": tid,
            "Theme_Label": theme_label,
            "Theme_Status": theme_status,
            "Central_Organizing_Concept": _txt(tr.get("Central_Organizing_Concept")),
            "Theme_Boundary": _txt(tr.get("Theme_Boundary")),
            "Closest_Competing_Theme": _txt(tr.get("Closest_Competing_Theme")),
            "Cluster_Count": len(cids),
        })

        for cid in cids:
            cr = cluster_lookup.get(cid)
            cluster_label = _txt(cr.get("Working_Cluster_Label")) if cr is not None else ""
            cluster_status = _txt(cr.get("Cluster_Status")) if cr is not None else ""

            tree_rows.append({
                "Level": 3,
                "Node_Type": "Cluster",
                "Node_ID": cid,
                "Parent_ID": tid,
                "Label": cluster_label,
                "Status": cluster_status,
                "Path": f"{did} > {tid} > {cid}",
            })

            members = cluster_members(frames, cid)
            if not members.empty and "Code_ID" in members.columns:
                members = members.drop_duplicates(subset=["Code_ID"], keep="first")

            cluster_rows.append({
                "Dimension_ID": did,
                "Theme_ID": tid,
                "Theme_Label": theme_label,
                "Cluster_ID": cid,
                "Cluster_Label": cluster_label,
                "Cluster_Status": cluster_status,
                "Core_Organizational_Function": _txt(cr.get("Core_Organizational_Function")) if cr is not None else "",
                "Operational_Definition": _txt(cr.get("Operational_Definition")) if cr is not None else "",
                "Inclusion_Boundary": _txt(cr.get("Inclusion_Boundary")) if cr is not None else "",
                "Exclusion_Boundary": _txt(cr.get("Exclusion_Boundary")) if cr is not None else "",
                "Nearest_Conceptual_Neighbours": _txt(cr.get("Nearest_Conceptual_Neighbours")) if cr is not None else "",
                "FOC_Count": int(members["Code_ID"].nunique()) if not members.empty and "Code_ID" in members.columns else 0,
                "Evidence_Count": int(members["Evidence_ID"].nunique()) if not members.empty and "Evidence_ID" in members.columns else 0,
                "Study_Count": int(members["Study_ID"].nunique()) if not members.empty and "Study_ID" in members.columns else 0,
            })

            if members.empty:
                continue

            for _, mr in members.iterrows():
                code_id = _txt(mr.get("Code_ID")).strip()
                evidence_id = _txt(mr.get("Evidence_ID")).strip()
                study_id = _txt(mr.get("Study_ID")).strip()
                foc = _txt(mr.get("First_Order_Code"))
                study_meta = _study_fields(studies.get(study_id, {}))

                trace_path = " > ".join(
                    x for x in [did, tid, cid, code_id, evidence_id, study_id] if x
                )

                row = {
                    "Trace_Path": trace_path,
                    "Dimension_ID": did,
                    "Candidate_Dimension_Name": dimension_name,
                    "Dimension_Status": dimension_status,
                    "Analytical_Definition": analytical_definition,
                    "Theme_ID": tid,
                    "Theme_Label": theme_label,
                    "Theme_Status": theme_status,
                    "Theme_Central_Concept": _txt(tr.get("Central_Organizing_Concept")),
                    "Cluster_ID": cid,
                    "Cluster_Label": cluster_label,
                    "Cluster_Status": cluster_status,
                    "Core_Organizational_Function": _txt(cr.get("Core_Organizational_Function")) if cr is not None else "",
                    "Operational_Definition": _txt(cr.get("Operational_Definition")) if cr is not None else "",
                    "Mapping_ID": _txt(mr.get("Mapping_ID")),
                    "Mapping_Status": _txt(mr.get("Mapping_Status")),
                    "Mapping_Rationale": _txt(mr.get("Mapping_Rationale")),
                    "Boundary_Rationale": _txt(mr.get("Boundary_Rationale")),
                    "Closest_Competing_Cluster": _txt(mr.get("Closest_Competing_Cluster")),
                    "Code_ID": code_id,
                    "First_Order_Code": foc,
                    "Code_Fidelity_Status": _txt(mr.get("Code_Fidelity_Status")),
                    "Source_to_Code_Rationale": _txt(mr.get("Source_to_Code_Rationale")),
                    "Evidence_ID": evidence_id,
                    "Evidence_Gate": _txt(mr.get("PM_Practice_Maturity_Evidence_Gate")),
                    "Meaning_Unit_Verbatim": _txt(mr.get("Meaning_Unit_Verbatim")),
                    "Context_Verbatim": _txt(mr.get("Context_Verbatim")),
                    "Original_Author_Term": _txt(mr.get("Original_Author_Term")),
                    "Author_Parent_Construct": _txt(mr.get("Author_Parent_Construct")),
                    "Author_Defined_Relationship": _txt(mr.get("Author_Defined_Relationship")),
                    "Printed_Page": _txt(mr.get("Printed_Page")),
                    "Study_ID": study_id,
                    **study_meta,
                }
                full_rows.append(row)

                tree_rows.append({
                    "Level": 4,
                    "Node_Type": "First-Order Code",
                    "Node_ID": code_id,
                    "Parent_ID": cid,
                    "Label": foc,
                    "Status": _txt(mr.get("Code_Fidelity_Status")),
                    "Path": f"{did} > {tid} > {cid} > {code_id}",
                })
                if evidence_id:
                    tree_rows.append({
                        "Level": 5,
                        "Node_Type": "Meaning Unit",
                        "Node_ID": evidence_id,
                        "Parent_ID": code_id,
                        "Label": _txt(mr.get("Meaning_Unit_Verbatim")),
                        "Status": _txt(mr.get("PM_Practice_Maturity_Evidence_Gate")),
                        "Path": f"{did} > {tid} > {cid} > {code_id} > {evidence_id}",
                    })
                if study_id:
                    tree_rows.append({
                        "Level": 6,
                        "Node_Type": "Study",
                        "Node_ID": study_id,
                        "Parent_ID": evidence_id or code_id,
                        "Label": study_meta.get("Study_Title", ""),
                        "Status": "",
                        "Path": trace_path,
                    })

    full_df = pd.DataFrame(full_rows)
    theme_df = pd.DataFrame(theme_rows)
    cluster_df = pd.DataFrame(cluster_rows)
    tree_df = pd.DataFrame(tree_rows).drop_duplicates(
        subset=["Node_Type", "Node_ID", "Parent_ID", "Path"],
        keep="first",
    )

    if full_df.empty:
        foc_df = pd.DataFrame()
        studies_df = pd.DataFrame()
    else:
        foc_cols = [
            "Dimension_ID", "Theme_ID", "Theme_Label", "Cluster_ID", "Cluster_Label",
            "Code_ID", "First_Order_Code", "Code_Fidelity_Status", "Mapping_ID",
            "Mapping_Status", "Mapping_Rationale", "Boundary_Rationale",
            "Closest_Competing_Cluster", "Evidence_ID", "Study_ID",
        ]
        foc_df = full_df[[c for c in foc_cols if c in full_df.columns]].drop_duplicates(
            subset=["Code_ID"], keep="first"
        )

        study_base = full_df[
            [
                "Study_ID", "Study_Title", "Authors", "Publication_Year", "DOI",
                "Country_Context", "Sector_Context", "Methodological_Family",
                "Publication_Type",
            ]
        ].drop_duplicates(subset=["Study_ID"], keep="first")
        metrics = (
            full_df.groupby("Study_ID", dropna=False)
            .agg(
                Theme_Count=("Theme_ID", "nunique"),
                Cluster_Count=("Cluster_ID", "nunique"),
                FOC_Count=("Code_ID", "nunique"),
                Evidence_Count=("Evidence_ID", "nunique"),
            )
            .reset_index()
        )
        studies_df = study_base.merge(metrics, on="Study_ID", how="left")

    if not theme_df.empty and not full_df.empty:
        tm = (
            full_df.groupby("Theme_ID")
            .agg(
                FOC_Count=("Code_ID", "nunique"),
                Evidence_Count=("Evidence_ID", "nunique"),
                Study_Count=("Study_ID", "nunique"),
            )
            .reset_index()
        )
        theme_df = theme_df.merge(tm, on="Theme_ID", how="left")
        for col in ["FOC_Count", "Evidence_Count", "Study_Count"]:
            theme_df[col] = theme_df[col].fillna(0).astype(int)

    snapshot = master_snapshot(frames)
    summary = {
        "Dimension_ID": did,
        "Candidate_Dimension_Name": dimension_name,
        "Dimension_Status": dimension_status,
        "Analytical_Definition": analytical_definition,
        "Theme_Count": len(theme_ids),
        "Cluster_Count": int(cluster_df["Cluster_ID"].nunique()) if not cluster_df.empty else 0,
        "FOC_Count": int(full_df["Code_ID"].nunique()) if not full_df.empty else 0,
        "Evidence_Count": int(full_df["Evidence_ID"].nunique()) if not full_df.empty else 0,
        "Study_Count": int(full_df["Study_ID"].nunique()) if not full_df.empty else 0,
        "Source_MASTER_Version": snapshot.get("version", ""),
        "Source_MASTER_Decision": snapshot.get("decision", ""),
    }

    return summary, {
        "01_Full_Trace": full_df,
        "02_Themes": theme_df,
        "03_Clusters": cluster_df,
        "04_FOCs": foc_df,
        "05_Studies": studies_df,
        "06_Tree_Index": tree_df,
    }


def _write_info_sheet(ws, summary: Dict[str, Any]) -> None:
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:D1")
    ws["A1"] = "PMM Dimension Traceability Export"
    ws["A1"].fill = TITLE_FILL
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=15)
    ws["A1"].alignment = Alignment(horizontal="left")

    ws.merge_cells("A2:D2")
    ws["A2"] = "Dimension → Theme → PCL/Cluster → FOC → Meaning Unit → Study"
    ws["A2"].fill = SUBTITLE_FILL
    ws["A2"].font = Font(bold=True, color="1F4E78")

    rows = [
        ("Dimension ID", summary.get("Dimension_ID", "")),
        ("Dimension Name", summary.get("Candidate_Dimension_Name", "")),
        ("Dimension Status", summary.get("Dimension_Status", "")),
        ("Analytical Definition", summary.get("Analytical_Definition", "")),
        ("Themes", summary.get("Theme_Count", 0)),
        ("Clusters / PCLs", summary.get("Cluster_Count", 0)),
        ("First-Order Codes", summary.get("FOC_Count", 0)),
        ("Meaning Units / Evidence", summary.get("Evidence_Count", 0)),
        ("Contributing Studies", summary.get("Study_Count", 0)),
        ("Source MASTER Version", summary.get("Source_MASTER_Version", "")),
        ("Source MASTER Decision", summary.get("Source_MASTER_Decision", "")),
        ("Export generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")),
    ]

    start = 4
    for i, (label, value) in enumerate(rows, start=start):
        ws.cell(i, 1, label)
        ws.cell(i, 2, _txt(value))
        ws.cell(i, 1).font = STATIC_FONT
        ws.cell(i, 2).font = LINKED_FONT
        ws.cell(i, 1).fill = PatternFill("solid", fgColor="F2F2F2")
        ws.cell(i, 1).alignment = Alignment(vertical="top")
        ws.cell(i, 2).alignment = Alignment(vertical="top", wrap_text=True)

    ws["A18"] = "How to use this workbook"
    ws["A18"].font = CONTROL_FONT
    ws["A18"].fill = CONTROL_FILL
    ws["A19"] = (
        "01_Full_Trace contains one complete trace row from the selected current dimension "
        "to the source-near Meaning Unit and contributing Study. Other sheets provide "
        "level-specific summaries. The export is read-only evidence for review; it does not modify the MASTER."
    )
    ws.merge_cells("A19:D21")
    ws["A19"].alignment = Alignment(wrap_text=True, vertical="top")

    links = [
        ("Full trace", "01_Full_Trace"),
        ("Themes", "02_Themes"),
        ("Clusters / PCLs", "03_Clusters"),
        ("First-Order Codes", "04_FOCs"),
        ("Studies", "05_Studies"),
        ("Tree index", "06_Tree_Index"),
    ]
    row = 23
    for label, sheet in links:
        ws.cell(row, 1, label)
        ws.cell(row, 2, f"='{sheet}'!A1")
        ws.cell(row, 2).hyperlink = f"#'{sheet}'!A1"
        ws.cell(row, 2).style = "Hyperlink"
        row += 1

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 80
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18


def _write_dataframe_sheet(ws, df: pd.DataFrame, sheet_name: str) -> None:
    ws.sheet_view.showGridLines = False
    if df is None or df.empty:
        ws["A1"] = "No rows available for this level in the selected current dimension."
        ws["A1"].fill = CAUTION_FILL
        ws["A1"].font = Font(bold=True, color="9C5700")
        return

    columns = [str(c) for c in df.columns]
    for col_idx, col in enumerate(columns, start=1):
        cell = ws.cell(1, col_idx, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GRAY)

    for r_idx, row in enumerate(df.itertuples(index=False, name=None), start=2):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(r_idx, c_idx, _txt(value))
            cell.font = LINKED_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32

    long_cols = {
        "Trace_Path", "Candidate_Dimension_Name", "Analytical_Definition",
        "Theme_Label", "Theme_Central_Concept", "Cluster_Label",
        "Core_Organizational_Function", "Operational_Definition",
        "Mapping_Rationale", "Boundary_Rationale", "First_Order_Code",
        "Source_to_Code_Rationale", "Meaning_Unit_Verbatim", "Context_Verbatim",
        "Original_Author_Term", "Author_Parent_Construct",
        "Author_Defined_Relationship", "Study_Title", "Label", "Path",
    }
    medium_cols = {
        "Closest_Competing_Cluster", "Country_Context", "Sector_Context",
        "Methodological_Family", "Authors", "DOI",
    }

    for idx, col in enumerate(columns, start=1):
        if col in long_cols:
            width = 42
        elif col in medium_cols:
            width = 28
        elif col.endswith("_ID") or col in {"Node_ID", "Parent_ID", "Level"}:
            width = 18
        elif "Status" in col or "Count" in col or col in {"Publication_Year", "Printed_Page"}:
            width = 16
        else:
            width = 22
        ws.column_dimensions[get_column_letter(idx)].width = width


def dimension_trace_workbook_bytes(
    frames: Dict[str, pd.DataFrame],
    dimension_id: str,
) -> Tuple[bytes, Dict[str, Any]]:
    summary, tables = build_dimension_trace_tables(frames, dimension_id)

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    info = wb.create_sheet("00_Export_Info")
    _write_info_sheet(info, summary)

    for sheet_name, df in tables.items():
        ws = wb.create_sheet(sheet_name[:31])
        _write_dataframe_sheet(ws, df, sheet_name)

    output = BytesIO()
    wb.save(output)
    return output.getvalue(), summary
